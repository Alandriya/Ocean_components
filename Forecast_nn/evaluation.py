"""Model-agnostic metrics, XLSX export and saved plotting samples."""
from pathlib import Path
import numpy as np
import torch


def _excel_value(value):
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, (tuple, list)):
        if len(value) == 2 and all(isinstance(v, (int, np.integer)) for v in value):
            return f'{int(value[0])}x{int(value[1])}'
        return ', '.join(str(_excel_value(v)) for v in value)
    if isinstance(value, dict):
        return ', '.join(f'{k}={_excel_value(v)}' for k, v in value.items())
    return value


def _short(value):
    """Compact, filesystem-safe number representation for experiment IDs."""
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, (float, np.floating)):
        value = float(value)
        decimal = f'{value:g}'
        if not np.isfinite(value) or value == 0:
            return decimal
        mantissa, exponent = f'{value:.6e}'.split('e')
        sci = f'{float(mantissa):g}e{int(exponent)}'
        return sci if len(sci) < len(decimal) else decimal
    return str(value)


def build_run_tag(model_name, variable, downsample, in_len, out_len, stride, model_params):
    """Short experiment name, relative to results/<model>/<variable>/.

    Always distinguish grid/history/horizon and the architecture's main size
    hyperparameters. Include secondary settings ONLY when different from the
    repository defaults. Unrecognized parameters are included verbatim to avoid
    silently giving distinct future architectures identical names.

    Split boundaries, dates and data file paths deliberately do not appear here;
    they are recorded in checkpoint metadata instead.
    """
    name = str(model_name).lower()
    params = dict(model_params)
    parts = [f'ds{int(downsample)}', f'i{int(in_len)}', f'o{int(out_len)}']
    if int(stride) != 1:
        parts.append(f's{int(stride)}')

    # Main architecture hyperparameters, included even when at default values.
    if name == 'swin_gru':
        core = [('base_channels', 'b'), ('temporal_tokens', 'tt'),
                ('attention_depth', 'ad'), ('heads', 'h'), ('window_size', 'ws')]
    elif name.startswith('fno'):
        core = [('width', 'w'), ('modes_y', 'my'), ('modes_x', 'mx'), ('layers', 'l')]
    elif name.startswith('simvp'):
        core = [('base_channels', 'b'), ('depth', 'd'), ('temporal_blocks', 'tb')]
    else:
        core = [('base_channels', 'b'), ('depth', 'd')]

    included = set()
    for key, short_name in core:
        if key in params:
            parts.append(f'{short_name}{_short(params[key])}')
            included.add(key)

    # These settings are real ablation knobs; omit the defaults to keep names short.
    defaults = {
        'dropout': 0.0 if name.startswith('fno') else 0.05,
        'residual': True,
        'use_mask_channel': True,
        'use_coord_channels': True,
        'mask_features': True,
        'lstm_kernel': 3,
        'temporal_kernel': 3,
        'spatial_kernel': 3,
        'horizon_embedding': True,
        'projection_width': 96,
        'padding': 2,
        'sde_condition_scale': 0.10,
        'sde_history': True,
    }
    aliases = {
        'dropout': 'p', 'residual': 'r', 'use_mask_channel': 'mc',
        'use_coord_channels': 'xy', 'mask_features': 'mf',
        'lstm_kernel': 'k', 'temporal_kernel': 'tk', 'spatial_kernel': 'sk',
        'horizon_embedding': 'he', 'projection_width': 'pw',
        'padding': 'pad', 'sde_condition_scale': 'sc', 'sde_history': 'sh',
    }
    for key, value in params.items():
        if key in included:
            continue
        if key in defaults:
            if value != defaults[key]:
                parts.append(f'{aliases[key]}{_short(value)}')
        else:
            # Do not hide future model settings from experiment identifiers.
            parts.append(f'{key}{_short(value)}')
    return '_'.join(parts)


def append_sde_run_tag(run_tag, sde_weight, b_kind, loss_mode=None):
    """Add only SDE loss/interpretation settings that change the experiment."""
    if float(sde_weight) > 0:
        run_tag += f'_sw{_short(float(sde_weight))}'
    if str(b_kind) != 'diffusion':
        run_tag += '_bv' if str(b_kind) == 'variance' else f'_b{b_kind}'
    if float(sde_weight) > 0 and loss_mode and str(loss_mode) != 'first_step':
        run_tag += '_ma' if str(loss_mode) == 'last_observed_all' else f'_m{loss_mode}'
    return run_tag


def add_run_id(run_tag, run_id=None):
    """Optional user-supplied suffix for runs varying training/split settings."""
    import re
    if not run_id:
        return run_tag
    run_id = str(run_id)
    if len(run_id) > 24 or not re.fullmatch(r'[A-Za-z0-9-]+', run_id):
        raise ValueError('--run-id must contain 1-24 ASCII letters, numbers or hyphens.')
    return f'{run_tag}_{run_id}'


def _gaussian_kernel(size=11, sigma=1.5):
    size = int(size)
    if size < 3 or size % 2 == 0:
        raise ValueError('SSIM window must be an odd integer >= 3.')
    x = np.arange(size, dtype=np.float64) - size // 2
    k = np.exp(-(x * x) / (2.0 * float(sigma) ** 2))
    return k / k.sum()


def _separable_blur_2d(arr, kernel):
    """Reflect-padded separable convolution without scipy dependency."""
    arr = np.asarray(arr, dtype=np.float64)
    pad = len(kernel) // 2
    h, w = arr.shape
    vp = np.pad(arr, ((pad, pad), (0, 0)), mode='reflect')
    tmp = np.zeros_like(arr, dtype=np.float64)
    for i, weight in enumerate(kernel):
        tmp += float(weight) * vp[i:i + h]
    hp = np.pad(tmp, ((0, 0), (pad, pad)), mode='reflect')
    out = np.zeros_like(arr, dtype=np.float64)
    for i, weight in enumerate(kernel):
        out += float(weight) * hp[:, i:i + w]
    return out


def masked_ssim2d(pred, truth, mask, window=11, sigma=1.5):
    """Masked local SSIM for one 2-D physical field.

    Local moments are computed with a Gaussian window using only ocean pixels.
    The dynamic range is derived from the truth map, which keeps the comparison
    consistent across NN and baseline predictions for the same target.
    """
    pred = np.asarray(pred, dtype=np.float64)
    truth = np.asarray(truth, dtype=np.float64)
    valid = np.asarray(mask, dtype=bool)
    if pred.shape != truth.shape or pred.shape != valid.shape:
        raise ValueError(f'SSIM shape mismatch: pred={pred.shape}, truth={truth.shape}, mask={valid.shape}.')
    if not np.any(valid):
        return float('nan')

    kernel = _gaussian_kernel(window, sigma)
    m = valid.astype(np.float64)
    support = _separable_blur_2d(m, kernel)
    denom = np.maximum(support, 1e-12)

    mu_x = _separable_blur_2d(pred * m, kernel) / denom
    mu_y = _separable_blur_2d(truth * m, kernel) / denom
    ex2 = _separable_blur_2d(pred * pred * m, kernel) / denom
    ey2 = _separable_blur_2d(truth * truth * m, kernel) / denom
    exy = _separable_blur_2d(pred * truth * m, kernel) / denom
    var_x = np.maximum(ex2 - mu_x * mu_x, 0.0)
    var_y = np.maximum(ey2 - mu_y * mu_y, 0.0)
    cov_xy = exy - mu_x * mu_y

    vals = truth[valid]
    data_range = float(np.nanmax(vals) - np.nanmin(vals))
    if not np.isfinite(data_range) or data_range < 1e-8:
        data_range = max(float(np.nanmax(np.abs(vals))), 1.0)
    c1 = (0.01 * data_range) ** 2
    c2 = (0.03 * data_range) ** 2
    num = (2.0 * mu_x * mu_y + c1) * (2.0 * cov_xy + c2)
    den = (mu_x * mu_x + mu_y * mu_y + c1) * (var_x + var_y + c2)
    ssim_map = num / np.maximum(den, 1e-30)

    # Require reasonable local ocean support to avoid coastline windows dominated
    # by reflected land/missing values.
    local_valid = valid & (support >= 0.5)
    if not np.any(local_valid):
        local_valid = valid
    return float(np.nanmean(ssim_map[local_valid]))


def masked_ssim(pred, truth, mask, window=11, sigma=1.5):
    """Average SSIM across channels for [C,H,W] fields."""
    pred = np.asarray(pred)
    truth = np.asarray(truth)
    if pred.ndim == 2:
        return masked_ssim2d(pred, truth, mask, window, sigma)
    scores = [masked_ssim2d(pred[c], truth[c], mask, window, sigma) for c in range(pred.shape[0])]
    return float(np.nanmean(scores))


def psnr_from_mse(mse, data_range):
    """10 log10(R^2/MSE) with fixed, positive R in physical units.

    Returns +inf for an exact (zero-MSE) prediction.
    """
    mse = np.asarray(mse, dtype=np.float64)
    data_range = np.asarray(data_range, dtype=np.float64)
    with np.errstate(divide='ignore', invalid='ignore'):
        return 10.0 * np.log10(np.square(data_range) / mse)


class OnlineErrorMetrics:
    """Efficient train/val MSE, RMSE, PSNR aggregation, no SSIM computation.

    Inputs are normalized tensors, but the tracked errors are physical (W/m2),
    and mask is always applied. Aggregation occurs *before* logarithms.
    """
    def __init__(self, out_len, std, psnr_ranges):
        self.std = np.asarray(std, dtype=np.float64).reshape(-1)
        self.ranges = np.asarray(psnr_ranges, dtype=np.float64).reshape(-1)
        self.sq_sum = np.zeros((int(out_len), len(self.std)), dtype=np.float64)
        self.count = np.zeros((int(out_len), len(self.std)), dtype=np.float64)

    @staticmethod
    def _valid(mask, x):
        if mask.ndim == 2:
            mask = mask[None, None, None]
        elif mask.ndim == 3:
            mask = mask[None, None]
        elif mask.ndim == 4:
            mask = mask[:, None]
        return mask.to(device=x.device, dtype=torch.bool).expand_as(x)

    def update(self, pred, target, mask):
        valid = self._valid(mask, pred)
        errors = (pred.detach().float() - target.detach().float()).square()
        total = (errors * valid).sum(dim=(0, 3, 4)).cpu().numpy()
        count = valid.sum(dim=(0, 3, 4)).cpu().numpy()
        self.sq_sum += total * self.std[None, :] ** 2
        self.count += count

    def summary(self):
        mse_hc = self.sq_sum / np.maximum(self.count, 1)
        psnr_hc = psnr_from_mse(mse_hc, self.ranges[None, :])
        overall_mse = float(self.sq_sum.sum() / max(self.count.sum(), 1))
        # Average dB across channels, each computed from aggregate channel MSE.
        mse_c = self.sq_sum.sum(axis=0) / np.maximum(self.count.sum(axis=0), 1)
        psnr_c = psnr_from_mse(mse_c, self.ranges)
        return {
            'overall_mse': overall_mse,
            'overall_rmse': float(np.sqrt(overall_mse)),
            'overall_psnr': float(np.mean(psnr_c)),
            'psnr': np.mean(psnr_hc, axis=1),
            'psnr_per_channel': psnr_c,
        }


class MetricAccumulator:
    """Masked value errors, SSIM and temporal-change diagnostics in physical units."""
    def __init__(self, out_len, ssim_window=11, ssim_sigma=1.5, psnr_ranges=None):
        self.out_len = int(out_len)
        self.ssim_window = int(ssim_window)
        self.ssim_sigma = float(ssim_sigma)
        self.psnr_ranges = None if psnr_ranges is None else np.asarray(psnr_ranges, dtype=np.float64).reshape(-1)
        if self.psnr_ranges is None or np.any(self.psnr_ranges <= 0):
            raise ValueError('MetricAccumulator requires positive train-derived psnr_ranges.')
        self.channel_sq_sum = np.zeros((out_len, len(self.psnr_ranges)), dtype=np.float64)
        self.channel_count = np.zeros((out_len, len(self.psnr_ranges)), dtype=np.int64)
        self.sq_sum = np.zeros(out_len, dtype=np.float64)
        self.abs_sum = np.zeros(out_len, dtype=np.float64)
        self.count = np.zeros(out_len, dtype=np.int64)
        self.ssim_sum = np.zeros(out_len, dtype=np.float64)
        self.ssim_count = np.zeros(out_len, dtype=np.int64)
        self.sample_rows = []
        self.delta_sq_sum = 0.0
        self.delta_abs_pred_sum = 0.0
        self.delta_abs_true_sum = 0.0
        self.delta_count = 0
        self.dx_sum = self.dy_sum = self.dxx_sum = self.dyy_sum = self.dxy_sum = 0.0

    def update(self, pred, truth, last_input, mask, sample_offset):
        valid = mask.astype(bool)
        error = pred - truth
        for b in range(error.shape[0]):
            for t in range(self.out_len):
                e = error[b, t, :, valid].reshape(-1)
                self.sq_sum[t] += np.square(e, dtype=np.float64).sum()
                self.abs_sum[t] += np.abs(e).sum(dtype=np.float64)
                self.count[t] += e.size
                err_channels = np.square(error[b, t][:, valid], dtype=np.float64)
                if err_channels.shape[0] != len(self.psnr_ranges):
                    raise ValueError('PSNR channel count differs from prediction channels.')
                self.channel_sq_sum[t] += err_channels.sum(axis=1)
                self.channel_count[t] += err_channels.shape[1]
                mse_channels = err_channels.mean(axis=1)
                psnr = float(np.mean(psnr_from_mse(mse_channels, self.psnr_ranges)))
                mse = float(np.mean(np.square(e, dtype=np.float64)))
                mae = float(np.mean(np.abs(e)))
                ssim = masked_ssim(
                    pred[b, t], truth[b, t], valid,
                    window=self.ssim_window, sigma=self.ssim_sigma,
                )
                if np.isfinite(ssim):
                    self.ssim_sum[t] += ssim
                    self.ssim_count[t] += 1
                self.sample_rows.append((sample_offset + b, t + 1, mse, mae, float(np.sqrt(mse)), ssim, psnr))

        pred_full = np.concatenate([last_input[:, None], pred], axis=1)
        truth_full = np.concatenate([last_input[:, None], truth], axis=1)
        dp = np.diff(pred_full, axis=1)[:, :, :, valid].reshape(-1).astype(np.float64)
        dt = np.diff(truth_full, axis=1)[:, :, :, valid].reshape(-1).astype(np.float64)
        self.delta_sq_sum += np.square(dp - dt).sum()
        self.delta_abs_pred_sum += np.abs(dp).sum()
        self.delta_abs_true_sum += np.abs(dt).sum()
        self.delta_count += dp.size
        self.dx_sum += dp.sum(); self.dy_sum += dt.sum()
        self.dxx_sum += np.square(dp).sum(); self.dyy_sum += np.square(dt).sum(); self.dxy_sum += (dp * dt).sum()

    def summary(self):
        mse = self.sq_sum / np.maximum(self.count, 1)
        mae = self.abs_sum / np.maximum(self.count, 1)
        ssim = self.ssim_sum / np.maximum(self.ssim_count, 1)
        overall_mse = float(self.sq_sum.sum() / max(self.count.sum(), 1))
        mse_hc = self.channel_sq_sum / np.maximum(self.channel_count, 1)
        psnr_hc = psnr_from_mse(mse_hc, self.psnr_ranges[None, :])
        mse_c = self.channel_sq_sum.sum(axis=0) / np.maximum(self.channel_count.sum(axis=0), 1)
        psnr_c = psnr_from_mse(mse_c, self.psnr_ranges)
        n = max(self.delta_count, 1)
        cov = self.dxy_sum - self.dx_sum * self.dy_sum / n
        var_x = self.dxx_sum - self.dx_sum ** 2 / n
        var_y = self.dyy_sum - self.dy_sum ** 2 / n
        corr = cov / np.sqrt(max(var_x * var_y, 1e-30))
        total_ssim_count = max(int(self.ssim_count.sum()), 1)
        return {
            'mse': mse,
            'mae': mae,
            'rmse': np.sqrt(mse),
            'ssim': ssim,
            'psnr': psnr_hc.mean(axis=1),
            'psnr_per_horizon_channel': psnr_hc,
            'psnr_per_channel': psnr_c,
            'psnr_ranges': self.psnr_ranges,
            'overall_psnr': float(psnr_c.mean()),
            'overall_mse': overall_mse,
            'overall_mae': float(self.abs_sum.sum() / max(self.count.sum(), 1)),
            'overall_rmse': float(np.sqrt(overall_mse)),
            'overall_ssim': float(self.ssim_sum.sum() / total_ssim_count),
            'delta_rmse': float(np.sqrt(self.delta_sq_sum / n)),
            'pred_mean_abs_change': float(self.delta_abs_pred_sum / n),
            'true_mean_abs_change': float(self.delta_abs_true_sum / n),
            'delta_correlation': float(corr),
        }


def save_metrics_xlsx(path, method, variable, run_tag, metrics, checkpoint, sample_rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ImportError('XLSX export requires openpyxl: pip install openpyxl') from exc

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'summary'
    fill = PatternFill('solid', fgColor='D9EAF7')
    ws.append(['Parameter', 'Value'])
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.fill = fill

    rows = [
        ('Variable', variable), ('Method', method), ('Run tag', run_tag),
        ('Overall MSE', metrics['overall_mse']), ('Overall MAE', metrics['overall_mae']),
        ('Overall RMSE', metrics['overall_rmse']), ('Overall SSIM', metrics['overall_ssim']),
        ('Overall PSNR (dB)', metrics['overall_psnr']), ('PSNR ranges', metrics['psnr_ranges']),
        ('Delta RMSE', metrics['delta_rmse']), ('Predicted mean |delta|', metrics['pred_mean_abs_change']),
        ('True mean |delta|', metrics['true_mean_abs_change']), ('Delta correlation', metrics['delta_correlation']),
    ]
    for key in (
        'model_name', 'downsample', 'grid_shape', 'in_len', 'out_len', 'stride',
        'train_end_idx', 'val_start_idx', 'test_start_idx', 'data_end_idx',
        'data_last_date', 'best_epoch', 'best_val_rmse', 'epochs',
        'learning_rate', 'dynamic_weight', 'sde_enabled',
        'sde_weight', 'sde_b_kind', 'sde_loss_mode', 'sde_time_offset', 'sde_dt', 'sde_source',
    ):
        if key in checkpoint:
            rows.append((key, checkpoint.get(key)))
    for key, value in checkpoint.get('model_kwargs', {}).items():
        rows.append((key, value))
    for row in rows:
        ws.append([_excel_value(v) for v in row])
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 42

    wh = wb.create_sheet('per_horizon')
    wh.append(['Horizon', 'MSE', 'MAE', 'RMSE', 'SSIM', 'PSNR (dB)'])
    for cell in wh[1]:
        cell.font = Font(bold=True); cell.fill = fill
    for t in range(len(metrics['mse'])):
        wh.append([
            t + 1, float(metrics['mse'][t]), float(metrics['mae'][t]),
            float(metrics['rmse'][t]), float(metrics['ssim'][t]), float(metrics['psnr'][t]),
        ])

    wc = wb.create_sheet('per_channel')
    wc.append(['Channel', 'PSNR range', 'Overall PSNR (dB)'])
    for c, (width, psnr) in enumerate(zip(metrics['psnr_ranges'], metrics['psnr_per_channel'])):
        wc.append([c, float(width), float(psnr)])
    wch = wb.create_sheet('per_horizon_channel')
    wch.append(['Horizon', 'Channel', 'PSNR (dB)'])
    for t, row in enumerate(metrics['psnr_per_horizon_channel']):
        for c, value in enumerate(row):
            wch.append([t + 1, c, float(value)])

    wsamp = wb.create_sheet('per_sample')
    wsamp.append(['Sample', 'Horizon', 'MSE', 'MAE', 'RMSE', 'SSIM', 'PSNR (dB)'])
    for cell in wsamp[1]:
        cell.font = Font(bold=True); cell.fill = fill
    for row in sample_rows:
        wsamp.append([_excel_value(v) for v in row])

    for sheet in (ws, wh, wc, wch, wsamp):
        sheet.freeze_panes = 'A2'
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical='center')
    wb.save(path)


def save_plot_samples(path, records, metadata):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {key: np.asarray(value) for key, value in records.items()}
    for key, value in metadata.items():
        payload[f'meta_{key}'] = np.asarray(value)
    np.savez_compressed(path, **payload)
