import os
from pathlib import Path
import numpy as np


def _fmt(value):
    if isinstance(value, float):
        return f'{value:g}'.replace('+', '')
    return str(value)


def _excel_value(value):
    """Convert checkpoint metadata to values that openpyxl can write into one cell."""
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, (tuple, list)):
        # Shapes are easier to read as 81x91; other short sequences become comma-separated text.
        if len(value) == 2 and all(isinstance(v, (int, np.integer)) for v in value):
            return f'{int(value[0])}x{int(value[1])}'
        return ', '.join(str(_excel_value(v)) for v in value)
    if isinstance(value, dict):
        return ', '.join(f'{k}={_excel_value(v)}' for k, v in value.items())
    return value


def build_run_tag(variable, checkpoint):
    """Build a compact filename tag from the parameters that materially define the trained model."""
    kw = checkpoint.get('model_kwargs', {})
    parts = [variable, f'ds{checkpoint.get("downsample", 1)}', f'in{checkpoint.get("in_len")}', f'out{checkpoint.get("out_len")}', f'st{checkpoint.get("stride", 1)}',
             f'bins{checkpoint.get("sde_bins")}', f'lam{_fmt(float(checkpoint.get("sde_lambda", 0.0)))}',
             f'base{kw.get("base_channels")}', f'lat{kw.get("latent_channels")}', f'kan{kw.get("kan_dim")}',
             f'rbf{kw.get("rbf_centers")}', f'g{_fmt(float(kw.get("rbf_gamma", 1.0)))}', f'drop{_fmt(float(kw.get("dropout", 0.0)))}']
    return '_'.join(parts)


class MetricAccumulator:
    """Accumulate masked error statistics per forecast horizon and per sample."""
    def __init__(self, out_len):
        self.out_len = out_len
        self.sq_sum = np.zeros(out_len, dtype=np.float64)
        self.abs_sum = np.zeros(out_len, dtype=np.float64)
        self.count = np.zeros(out_len, dtype=np.int64)
        self.sample_rows = []

    def update(self, error, mask, sample_offset):
        # error: B,T,C,H,W in original physical units.
        valid = mask.astype(bool)
        for b in range(error.shape[0]):
            for t in range(self.out_len):
                e = error[b, t, :, valid].reshape(-1)
                self.sq_sum[t] += np.square(e, dtype=np.float64).sum()
                self.abs_sum[t] += np.abs(e).sum(dtype=np.float64)
                self.count[t] += e.size
                mse = float(np.mean(np.square(e, dtype=np.float64)))
                mae = float(np.mean(np.abs(e)))
                self.sample_rows.append((sample_offset + b, t + 1, mse, mae, float(np.sqrt(mse))))

    def summary(self):
        mse = self.sq_sum / np.maximum(self.count, 1)
        mae = self.abs_sum / np.maximum(self.count, 1)
        rmse = np.sqrt(mse)
        overall_mse = float(self.sq_sum.sum() / max(self.count.sum(), 1))
        overall_mae = float(self.abs_sum.sum() / max(self.count.sum(), 1))
        return {'mse': mse, 'mae': mae, 'rmse': rmse, 'overall_mse': overall_mse, 'overall_mae': overall_mae, 'overall_rmse': float(np.sqrt(overall_mse))}


def save_metrics_xlsx(path, method, variable, run_tag, metrics, checkpoint, sample_rows):
    """Save one forecast method to its own XLSX workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise ImportError('XLSX export requires openpyxl: pip install openpyxl') from exc

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'summary'
    header_fill = PatternFill('solid', fgColor='D9EAF7')
    rows = [('Variable', variable), ('Method', method), ('Run tag', run_tag), ('Overall MSE', metrics['overall_mse']), ('Overall MAE', metrics['overall_mae']), ('Overall RMSE', metrics['overall_rmse'])]
    for key in ('downsample', 'grid_shape', 'in_len', 'out_len', 'stride', 'sde_bins', 'sde_lambda', 'train_end_idx'):
        rows.append((key, checkpoint.get(key)))
    for key, value in checkpoint.get('model_kwargs', {}).items():
        rows.append((key, value))
    ws.append(['Parameter', 'Value'])
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.fill = header_fill
    for row in rows:
        ws.append([_excel_value(value) for value in row])
    ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 32

    wh = wb.create_sheet('per_horizon')
    wh.append(['Horizon', 'MSE', 'MAE', 'RMSE'])
    for cell in wh[1]:
        cell.font = Font(bold=True); cell.fill = header_fill
    for t in range(len(metrics['mse'])):
        wh.append([t + 1, float(metrics['mse'][t]), float(metrics['mae'][t]), float(metrics['rmse'][t])])
    for col, width in zip(('A', 'B', 'C', 'D'), (12, 18, 18, 18)):
        wh.column_dimensions[col].width = width

    wsamp = wb.create_sheet('per_sample')
    wsamp.append(['Sample', 'Horizon', 'MSE', 'MAE', 'RMSE'])
    for cell in wsamp[1]:
        cell.font = Font(bold=True); cell.fill = header_fill
    for row in sample_rows:
        wsamp.append([_excel_value(value) for value in row])
    for col, width in zip(('A', 'B', 'C', 'D', 'E'), (12, 12, 18, 18, 18)):
        wsamp.column_dimensions[col].width = width
    for sheet in (ws, wh, wsamp):
        sheet.freeze_panes = 'A2'
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical='center')
    wb.save(path)


def save_plot_samples(path, records, metadata):
    """Store selected forecasts and simple scalar metadata for the standalone plotting script."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    payload = {'metadata': np.array([str(metadata)], dtype=object)}
    for key, value in records.items():
        payload[key] = np.asarray(value)
    # Named metadata fields avoid parsing the legacy string representation in new plotting code.
    for key, value in metadata.items():
        payload[f'meta_{key}'] = np.asarray(value)
    np.savez_compressed(path, **payload)
